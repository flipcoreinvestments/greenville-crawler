#!/usr/bin/env python3
"""
SKIP-TRACE RESULTS INTAKE (added 2026-10-11, T Dawg's ask).

After she skip traces a list, she uploads the result file(s) to the private
Supabase Storage bucket "skiptrace-results":

    skiptrace-results/traced/       leads that came back with phone numbers
    skiptrace-results/not_traced/   leads that came back with nothing
    skiptrace-results/  (top level) a whole BatchData result file -- each
                                    row is sorted by whether it has a phone

Accepts BatchData's own result export (Input Data:* / Skiptrace:* columns),
our call-list CSV, or her GHL upload file (Street Address / phone columns),
as .csv or .xlsx. Every file is read once (remembered in skiptrace_files).

For every row it finds the lead (by property address + zip) and records:
    skiptrace_status   'traced' | 'failed'
    skiptrace_date     when the file was read
    skiptrace_input    the name/address that was actually searched, so the
                       second QA check (skiptrace_followup.py) can tell
                       whether a failed lead has since been fixed
    skiptrace_attempts how many times it has been traced
Rows that match no lead are listed in lead-exports/skiptrace_unmatched.csv.
Once a lead is marked, it never shows up on the morning call lists again.
"""
import csv
import io
import json
import os
import re
import sys
import tempfile

import psycopg2

from lead_common import ensure_schema
from skiptrace_prep import compute_skiptrace_identity

SUPABASE_URL = os.environ.get("SUPABASE_URL", "https://izxqrskybwtqzbbbzlux.supabase.co")
BUCKET = "skiptrace-results"
EXPORT_BUCKET = "lead-exports"

ADDR_COLS = ["Input Data:Property Address", "Input Data:Street Address", "Property Address",
             "Street Address", "Address", "Skiptrace Address"]
CITY_COLS = ["Input Data:Property City", "Input Data:City", "Property City", "City"]
ZIP_COLS = ["Input Data:Property Zip", "Input Data:Postal Code", "Property Zip", "Postal Code", "Zip"]
FIRST_COLS = ["Input Data:first name", "Input Data:First Name", "Skiptrace First Name", "first name", "First Name"]
LAST_COLS = ["Input Data:last name", "Input Data:Last Name", "Skiptrace Last Name", "last name", "Last Name"]
# 2026-10-16: the address that was actually searched (Skiptrace Address on
# the pull list; "Mailing Address" in BatchData's mapped files), kept so a
# retrace that would search the same name AND address again is caught.
SEARCH_ADDR_COLS = ["Input Data:Skiptrace Address", "Input Data:Mailing Address", "Skiptrace Address",
                    "Mailing Address"]
PHONE_RE = re.compile(r"(phonenumbers\.\d+\.number|^phone( number \d+)?$|^phone\s*\d*$)", re.I)


def ensure_columns(conn):
    with conn.cursor() as cur:
        for col, typ in [("skiptrace_status", "text"), ("skiptrace_date", "timestamptz"),
                         ("skiptrace_input", "jsonb"), ("skiptrace_attempts", "int"),
                         ("skiptrace_fix_status", "text"), ("skiptrace_fix_note", "text"),
                         ("skiptrace_snapshot", "jsonb"), ("active_months", "text[]"),
                         # 2026-10-13: the phones a trace returned (private DB only), so
                         # dnc_scrub.py can check them and text_ready.csv can list them
                         ("skiptrace_phones", "text[]"),
                         # 2026-10-16: GHL sync (ghl_sync.py)
                         ("ghl_phone", "text"), ("ghl_in_crm", "boolean"), ("retrace_reason", "text"),
                         ("ghl_contact_id", "text"), ("ghl_seen_at", "timestamptz"), ("ghl_tags", "text[]"),
                         ("ghl_reset_at", "timestamptz")]:
            cur.execute(f"alter table leads add column if not exists {col} {typ}")
        cur.execute("""create table if not exists skiptrace_files (
                         name text primary key, rows_read int, matched int, traced int, failed int,
                         read_at timestamptz not null default now())""")
        cur.execute("""create table if not exists lead_error_log (
                         id bigserial primary key, lead_id uuid not null, found_at timestamptz not null default now(),
                         reason text not null)""")
        cur.execute("create index if not exists idx_lead_error_log_found on lead_error_log(found_at)")


def _first(row, cols):
    lower = {k.lower().strip(): v for k, v in row.items() if k}
    for c in cols:
        v = lower.get(c.lower())
        if v not in (None, ""):
            return str(v).strip()
    return ""


def row_phones(row):
    """Every 10-digit phone in a result row, in column order, no repeats."""
    out = []
    for k, v in row.items():
        if PHONE_RE.search((k or "").strip()):
            d = re.sub(r"\D", "", str(v or ""))
            d = d[1:] if len(d) == 11 and d.startswith("1") else d
            if len(d) == 10 and d not in out:
                out.append(d)
    return out


def has_phone(row):
    return any(PHONE_RE.search((k or "").strip()) and re.sub(r"\D", "", str(v or "")).__len__() >= 10
               for k, v in row.items())


def read_rows(path):
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [dict(r) for r in csv.DictReader(f)]
    from openpyxl import load_workbook
    ws = load_workbook(path, data_only=True).active
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    return [{h: ("" if v is None else str(v)) for h, v in zip(header, r) if h}
            for r in it if any(c is not None for c in r)]


def list_files(session, key, url=SUPABASE_URL):
    h = {"Authorization": f"Bearer {key}", "apikey": key}
    # create the private bucket the first time (409/400 = already exists)
    session.post(f"{url}/storage/v1/bucket", headers=h,
                 json={"id": BUCKET, "name": BUCKET, "public": False}, timeout=60)
    out = []
    for prefix in ("", "traced/", "not_traced/"):
        r = session.post(f"{url}/storage/v1/object/list/{BUCKET}", headers=h,
                         json={"prefix": prefix, "limit": 1000, "sortBy": {"column": "name", "order": "asc"}},
                         timeout=60)
        if r.status_code != 200:
            continue
        for obj in r.json():
            name = obj.get("name") or ""
            if obj.get("id") and name.lower().endswith((".csv", ".xlsx")):
                out.append((prefix + name, obj.get("updated_at") or ""))
    return out


def _same_person(a_first, a_last, b_first, b_last):
    """Same rule as export_call_lists._same_person: same last name and the
    same first 3 letters of the first name."""
    n = lambda x: re.sub(r"[^a-z]", "", (x or "").lower())
    if n(a_last) != n(b_last):
        return False
    fa, fb = n(a_first), n(b_first)
    return not fa or not fb or fa[:3] == fb[:3]


def mark_row(cur, row, status, when_label):
    # old exports put the '*'/'***' flag in front of the address
    addr = re.sub(r"^[^0-9A-Za-z]+", "", _first(row, ADDR_COLS))
    if not addr:
        return None
    z = re.sub(r"\D", "", _first(row, ZIP_COLS))[:5]
    cur.execute(
        """select id from leads
           where address_core(address) = address_core(%s)
             and (%s = '' or coalesce(zip, '') = '' or zip = %s)
           order by is_duplicate, is_sold, list_count desc limit 1""",
        (addr, z, z))
    hit = cur.fetchone()
    if not hit:
        return None
    inp = {"first": _first(row, FIRST_COLS), "last": _first(row, LAST_COLS), "address": addr,
           "city": _first(row, CITY_COLS), "zip": z, "file": when_label,
           "searched_address": _first(row, SEARCH_ADDR_COLS)}
    # 2026-10-15: a retrace (the person to call changed) replaces the old
    # person's phones instead of keeping them behind the new ones, and an
    # empty retrace shows as not traced rather than keeping the wrong phones.
    cur.execute("""select skiptrace_input, retrace_reason, owner_name, address, mailing_address, raw, source_tags
                   from leads where id = %s""", (hit[0],))
    old, was_wrong, owner, paddr, maddr, raw, stags = (cur.fetchone() or [None] * 7)
    old = old or {}
    # 2026-10-16: who the SYSTEM would have traced at this moment. When she
    # edited the name before tracing (an ORANGE row), the traced name differs
    # from the system's -- that is not a "person changed" retrace; only a later
    # change of the system's own person (a PR appointed...) is.
    try:
        sf, sl, _, _ = compute_skiptrace_identity(
            {"owner_name": owner, "address": paddr, "mailing_address": maddr, "raw": raw or {}}, stags or [])
        inp["system_first"], inp["system_last"] = sf or "", sl or ""
    except Exception:
        pass
    # ...and so does a new trace after a caller reported "wrong person" (even
    # with the same name -- she changed the address or used another provider):
    # the old contact was tagged retrace and deleted, its phones were wrong.
    if was_wrong or ((old.get("first") or old.get("last")) and (inp["first"] or inp["last"])
                     and not _same_person(old.get("first"), old.get("last"), inp["first"], inp["last"])):
        # 2026-10-16: the new person gets a NEW GHL contact (the old one was
        # tagged "retrace" and deleted), so the old contact's phone and the
        # retrace reason are forgotten too (ghl_sync.py columns).
        cur.execute("""update leads set skiptrace_phones = null, skiptrace_status = null,
                              ghl_phone = null, ghl_in_crm = null, ghl_contact_id = null, ghl_seen_at = null,
                              ghl_tags = null, retrace_reason = null, ghl_reset_at = now() where id = %s""",
                    (hit[0],))
    # a lead already marked traced stays traced (a later empty re-run of the
    # same lead never erases phones she already has)
    cur.execute(
        """update leads set
               skiptrace_status = case when skiptrace_status = 'traced' then 'traced' else %s end,
               skiptrace_date = now(),
               skiptrace_input = %s::jsonb,
               skiptrace_attempts = coalesce(skiptrace_attempts, 0) + 1,
               skiptrace_fix_status = null, skiptrace_fix_note = null
           where id = %s""",
        (status, json.dumps(inp), hit[0]))
    save_phones(cur, hit[0], row_phones(row))
    return hit[0]


def save_phones(cur, lead_id, phones):
    """Newest file's phones first, older ones kept after them (max 10)."""
    if not phones:
        return
    cur.execute("select coalesce(skiptrace_phones, '{}') from leads where id = %s", (lead_id,))
    old = list(cur.fetchone()[0] or [])
    merged = []
    for p in list(phones) + old:
        if p not in merged:
            merged.append(p)
    cur.execute("update leads set skiptrace_phones = %s where id = %s", (merged[:10], lead_id))


def backfill_phones(conn, key):
    """One time (2026-10-13): files read before phones were saved are read
    again for their phones only -- nothing else on the lead changes."""
    with conn.cursor() as cur:
        cur.execute("select 1 from pipeline_migrations where name = 'skiptrace_phones_backfill_v1'")
        if cur.fetchone():
            return 0
    import requests
    s = requests.Session()
    h = {"Authorization": f"Bearer {key}", "apikey": key}
    n = 0
    for name, _ in list_files(s, key):
        r = s.get(f"{SUPABASE_URL}/storage/v1/object/{BUCKET}/{name}", headers=h, timeout=120)
        if r.status_code != 200:
            continue
        with tempfile.NamedTemporaryFile(suffix=os.path.splitext(name)[1]) as tmp:
            tmp.write(r.content)
            tmp.flush()
            rows = read_rows(tmp.name)
        with conn.cursor() as cur:
            for row in rows:
                phones = row_phones(row)
                addr = re.sub(r"^[^0-9A-Za-z]+", "", _first(row, ADDR_COLS))
                if not phones or not addr:
                    continue
                z = re.sub(r"\D", "", _first(row, ZIP_COLS))[:5]
                cur.execute("""select id from leads where address_core(address) = address_core(%s)
                                 and (%s = '' or coalesce(zip, '') = '' or zip = %s)
                               order by is_duplicate, is_sold, list_count desc limit 1""", (addr, z, z))
                hit = cur.fetchone()
                if hit:
                    save_phones(cur, hit[0], phones)
                    n += 1
        conn.commit()
    with conn.cursor() as cur:
        cur.execute("insert into pipeline_migrations (name) values ('skiptrace_phones_backfill_v1') "
                    "on conflict do nothing")
    conn.commit()
    print(f"  phones backfilled from earlier result files: {n} rows")
    return n


def intake(conn, key):
    import requests
    s = requests.Session()
    h = {"Authorization": f"Bearer {key}", "apikey": key}
    with conn.cursor() as cur:
        cur.execute("select name from skiptrace_files")
        done = {r[0] for r in cur.fetchall()}
    unmatched = []
    totals = {"files": 0, "rows": 0, "matched": 0, "traced": 0, "failed": 0}
    for name, _ in list_files(s, key):
        if name in done:
            continue
        r = s.get(f"{SUPABASE_URL}/storage/v1/object/{BUCKET}/{name}", headers=h, timeout=120)
        r.raise_for_status()
        with tempfile.NamedTemporaryFile(suffix=os.path.splitext(name)[1]) as tmp:
            tmp.write(r.content)
            tmp.flush()
            rows = read_rows(tmp.name)
        forced = "traced" if name.startswith("traced/") else "failed" if name.startswith("not_traced/") else None
        m = t = f = 0
        with conn.cursor() as cur:
            for row in rows:
                status = forced or ("traced" if has_phone(row) else "failed")
                lead_id = mark_row(cur, row, status, name)
                if lead_id is None:
                    unmatched.append([name, _first(row, ADDR_COLS), _first(row, ZIP_COLS),
                                      _first(row, FIRST_COLS), _first(row, LAST_COLS)])
                    continue
                m += 1
                t += status == "traced"
                f += status == "failed"
            cur.execute("insert into skiptrace_files (name, rows_read, matched, traced, failed) "
                        "values (%s,%s,%s,%s,%s) on conflict (name) do nothing", (name, len(rows), m, t, f))
        conn.commit()
        print(f"  {name}: {len(rows)} rows, {m} matched ({t} traced, {f} not traced), "
              f"{len(rows) - m} unmatched")
        totals["files"] += 1
        totals["rows"] += len(rows)
        totals["matched"] += m
        totals["traced"] += t
        totals["failed"] += f
    if unmatched:
        buf = io.StringIO()
        w = csv.writer(buf)
        w.writerow(["File", "Address", "Zip", "First Name", "Last Name"])
        w.writerows(unmatched)
        s.post(f"{SUPABASE_URL}/storage/v1/object/{EXPORT_BUCKET}/skiptrace_unmatched.csv",
               headers={**h, "x-upsert": "true", "Content-Type": "text/csv"},
               data=buf.getvalue().encode("utf-8"), timeout=120)
    return totals, len(unmatched)


def main():
    db_url = os.environ.get("DATABASE_URL")
    key = os.environ.get("SUPABASE_SERVICE_KEY")
    if not db_url:
        print("DATABASE_URL is not set", file=sys.stderr)
        sys.exit(1)
    conn = psycopg2.connect(db_url)
    ensure_columns(conn)
    ensure_schema(conn)
    conn.commit()
    if not key:
        print("SUPABASE_SERVICE_KEY is not set -- can't read the skiptrace-results bucket")
        return
    backfill_phones(conn, key)
    totals, unmatched = intake(conn, key)
    with conn.cursor() as cur:
        cur.execute("insert into source_runs (source_name, records_found, records_new, notes) "
                    "values ('skiptrace_intake', %s, %s, %s)",
                    (totals["rows"], totals["matched"],
                     f"ok: files:{totals['files']} traced:{totals['traced']} not_traced:{totals['failed']} "
                     f"unmatched:{unmatched}"))
    conn.commit()
    conn.close()
    print(f"Skip-trace intake: {totals['files']} new file(s), {totals['matched']} lead(s) marked "
          f"({totals['traced']} traced, {totals['failed']} not traced), {unmatched} row(s) unmatched")


if __name__ == "__main__":
    main()
