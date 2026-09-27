#!/usr/bin/env python3
"""
PropStream export importer (added 2026-09-27).

T Dawg exports these PropStream lead lists for Greenville County, SC and
uploads the .xlsx/.csv files to the repo folder `propstream_exports/`:

  file name contains   ->  what it does
  "probate"            ->  tag pre_probate      (owner on title flagged deceased)
  "foreclosure"        ->  tag pre_foreclosure  (default/lis pendens recorded, last 6 months)
  "lien"               ->  tag involuntary_lien (HOA / mechanic's / utility / child support, last 2 yrs)
  "fail"               ->  tag failed_listing   (MLS failed in the last 12 months)
  "vacant"             ->  is_vacant = true     (USPS: no mail collected 90+ days)
                           Vacant is NOT a list: it never creates a lead or
                           counts as a list, it only strengthens a lead that
                           is already on one (and turns on the '***' flag
                           when that lead is also on the tax sale list).
  "assumable"          ->  has_assumable_loan = true (all open loans FHA/VA/USDA;
                           creative-finance flag, not a distress list)

Lists deliberately NOT imported (researched 2026-09-27): Divorce (1 record
in the county, recorder-sourced, months late), Bankruptcy (2 records),
Tax Delinquency / Auctions (county sources are fresher), owner-profile lists
(High Equity, Free & Clear, Senior, Tired Landlord, Absentee), Upside Down
(estimated balances), Bank Owned, Cash Buyers / Flippers (buyers, not sellers).

RULES (same standards as the county scripts):
  - Greenville County only.
  - Match on the parcel number first: PropStream's APN with punctuation
    removed is exactly the county PIN stored in leads.pin
    ("0624.01-05-003.00" -> "0624010500300"). Fall back to street address
    + zip. A property with no street number is only used if the PIN matches.
  - Skip anything on the market right now (Active / Under Contract /
    Pending / Contingent / Coming Soon) and anything SOLD in the last 12
    months -- 4 of the first 68 pre-probate rows were one or the other even
    with PropStream's Off Market filter on.
  - Residential only: the county's land-use code decides when we have it;
    otherwise PropStream's property type must be residential. (The first
    Vacant export had 700+ stores, offices, warehouses, churches.)
  - Business owners are NOT dropped here: they land on the separate
    business_owned_leads list by the same rule as every other source.
  - Freshness: only the NEWEST file per list counts. A property that drops
    off the newest export loses the tag. If the newest file for a list is
    more than STALE_DAYS old, that whole list is expired (PropStream data
    goes stale; T Dawg re-exports monthly).
"""

import csv
import glob
import json
import os
import re
import sys
from datetime import date, datetime, timedelta, timezone

import psycopg2

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from lead_common import (address_core, ensure_schema, expire_by_raw_key,  # noqa: E402
                         is_residential, remove_tag, rescore_all)

EXPORT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "propstream_exports")
STALE_DAYS = 60
SOLD_WINDOW_DAYS = 365

LISTS = [  # (filename keyword, kind, tag-or-column)
    ("probate", "tag", "pre_probate"),
    ("foreclosure", "tag", "pre_foreclosure"),
    ("lien", "tag", "involuntary_lien"),
    ("fail", "tag", "failed_listing"),
    ("vacant", "flag", "is_vacant"),
    ("assumable", "flag", "has_assumable_loan"),
]

ON_MARKET = {"ACTIVE", "ACTIVE UNDER CONTRACT", "PENDING", "CONTINGENT", "COMING SOON"}
PS_RESIDENTIAL = {
    "single family residential", "residential-vacant land", "mobile home", "manufactured home",
    "duplex (2 units, any combination)", "triplex (3 units, any combination)",
    "quadruplex (4 units, any combination)", "townhouse (residential)", "condominium (residential)",
    "residential (general) (single)", "residential (general)", "cluster home", "patio home",
}


# ------------------------------------------------------------------ reading
def _cell(v):
    if v is None:
        return None
    if isinstance(v, float) and v != v:  # NaN
        return None
    if isinstance(v, (datetime, date)):
        return v
    s = str(v).strip()
    return s or None


def read_rows(path):
    """Rows as dicts keyed by PropStream's column headers (.xlsx or .csv)."""
    if path.lower().endswith(".csv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            return [{k.strip(): _cell(v) for k, v in r.items() if k} for r in csv.DictReader(f)]
    from openpyxl import load_workbook
    # read_only=False: PropStream workbooks carry no dimension record, and
    # read-only mode then sees only the header row
    ws = load_workbook(path, data_only=True).active
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(it)]
    return [{h: _cell(v) for h, v in zip(header, r) if h} for r in it if any(c is not None for c in r)]


def file_date(path):
    """PropStream names exports like Property_Export_vacant9-27-26.xlsx."""
    m = re.search(r"(\d{1,2})-(\d{1,2})-(\d{2,4})(?=\D*$)", os.path.basename(path))
    if m:
        mo, d, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y + 2000 if y < 100 else y, mo, d)
        except ValueError:
            pass
    return None


def newest_files(export_dir):
    """{keyword: path} -- newest export per list, by the date in the file name."""
    files = [p for p in glob.glob(os.path.join(export_dir, "*"))
             if p.lower().endswith((".xlsx", ".csv"))]
    best = {}
    for kw, _, _ in LISTS:
        mine = [p for p in files if kw in os.path.basename(p).lower()]
        if mine:
            best[kw] = max(mine, key=lambda p: (file_date(p) or date.min, os.path.basename(p)))
    return best


def _as_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%Y-%m-%d %H:%M:%S"):
            try:
                return datetime.strptime(v[:19], fmt).date()
            except ValueError:
                continue
    return None


def pin_from_apn(apn):
    digits = re.sub(r"\D", "", str(apn or ""))
    return digits or None


def street_line(r):
    addr = r.get("Address")
    unit = r.get("Unit #")
    if not addr:
        return None
    return f"{addr} Unit {unit}" if unit else addr


def owner_county_style(r):
    """County writes owners LAST FIRST -- keep the same shape for name matching."""
    parts = []
    for n in ("1", "2"):
        last, first = r.get(f"Owner {n} Last Name"), r.get(f"Owner {n} First Name")
        if last or first:
            parts.append(" ".join(p for p in (last, first) if p))
    return " ".join(parts) or None


def screen(r, today):
    """None if the row is usable, else the reason it is skipped."""
    if (r.get("County") or "").strip().lower() != "greenville":
        return "not_greenville"
    status = (r.get("MLS Status") or "").strip().upper()
    if status in ON_MARKET:
        return "on_market"
    if status == "SOLD":
        d = _as_date(r.get("MLS Date"))
        if d is None or (today - d).days <= SOLD_WINDOW_DAYS:
            return "sold_recently"
    if not pin_from_apn(r.get("APN")):
        return "no_apn"
    return None


def ps_residential(r):
    return (r.get("Property Type") or "").strip().lower() in PS_RESIDENTIAL


# ------------------------------------------------------------------ database
def ensure_columns(conn):
    with conn.cursor() as cur:
        cur.execute("alter table leads add column if not exists has_assumable_loan boolean")


def find_lead(cur, pin, street, zip_code):
    cur.execute("select id, land_use from leads where pin = %s and is_duplicate = false "
                "order by is_sold, created_at limit 1", (pin,))
    hit = cur.fetchone()
    if hit or not street or not re.match(r"^\d", street):
        return hit
    cur.execute("select id, land_use from leads where address_core(address) = address_core(%s) "
                "and coalesce(zip, '') = %s and is_duplicate = false order by is_sold, created_at limit 2",
                (street, zip_code or ""))
    rows = cur.fetchall()
    return rows[0] if len(rows) == 1 else None  # ambiguous address -> don't guess


def payload(r, path, fdate):
    return {
        "pin": pin_from_apn(r.get("APN")),
        "file": os.path.basename(path),
        "file_date": fdate.isoformat() if fdate else None,
        "property_type": r.get("Property Type"),
        "mls_status": r.get("MLS Status"),
        "owner": owner_county_style(r),
        "mailing": " ".join(str(r.get(k)) for k in ("Mailing Address", "Mailing City", "Mailing State", "Mailing Zip")
                            if r.get(k)),
        "do_not_mail": r.get("Do Not Mail"),
        "est_equity": r.get("Est. Equity"),
        "foreclosure_factor": r.get("Foreclosure Factor"),
        "lien_amount": r.get("Lien Amount"),
        "imported_at": datetime.now(timezone.utc).isoformat(),
    }


def import_list(conn, kw, kind, target, path, today):
    fdate = file_date(path)
    rows = read_rows(path)
    raw_key = f"propstream_{target}" if kind == "tag" else f"propstream_{kw}"
    stats = {"rows": len(rows), "tagged": 0, "created": 0, "skipped": {}}
    kept_pins = []
    with conn.cursor() as cur:
        for r in rows:
            why = screen(r, today)
            pin = pin_from_apn(r.get("APN"))
            street = street_line(r)
            zip_code = (str(r.get("Zip") or "").split(".")[0])[:5] or None
            hit = None if why else find_lead(cur, pin, street, zip_code)
            if not why:
                land_use = hit[1] if hit else None
                residential = is_residential(land_use) if land_use else ps_residential(r)
                if not residential:
                    why = "commercial"
            if why:
                stats["skipped"][why] = stats["skipped"].get(why, 0) + 1
                continue
            body = json.dumps({raw_key: payload(r, path, fdate)})
            if hit:
                if kind == "tag":
                    cur.execute(
                        "update leads set source_tags = array(select distinct unnest(source_tags || array[%s])), "
                        "raw = coalesce(raw, '{}'::jsonb) || %s::jsonb, pin = coalesce(pin, %s), updated_at = now() "
                        "where id = %s", (target, body, pin, hit[0]))
                else:
                    cur.execute(
                        f"update leads set {target} = true, raw = coalesce(raw, '{{}}'::jsonb) || %s::jsonb, "
                        "updated_at = now() where id = %s", (body, hit[0]))
                stats["tagged"] += 1
                kept_pins.append(pin)
            elif kind == "tag" and street and re.match(r"^\d", street):
                # a real distress list with a real street address the county
                # data hasn't given us yet -> new lead (flagged by the usual
                # review rules until land use backfills)
                cur.execute(
                    """
                    insert into leads (address, city, state, zip, county, owner_name, mailing_address,
                                       source_tags, raw, pin)
                    values (%s, %s, 'SC', %s, 'Greenville', %s, %s, array[%s]::text[], %s::jsonb, %s)
                    on conflict (lower(address)) do update set
                        source_tags = array(select distinct unnest(leads.source_tags || excluded.source_tags)),
                        raw = coalesce(leads.raw, '{}'::jsonb) || excluded.raw,
                        pin = coalesce(leads.pin, excluded.pin), updated_at = now()
                    """,
                    (street, r.get("City"), zip_code, owner_county_style(r), r.get("Mailing Address"),
                     target, body, pin))
                stats["created"] += 1
                kept_pins.append(pin)
            else:
                stats["skipped"]["no_match"] = stats["skipped"].get("no_match", 0) + 1
    # expiry: anything tagged from an older export that isn't in this one
    if kind == "tag":
        expire_by_raw_key(conn, target, raw_key, "pin", kept_pins)
    else:
        expire_flag(conn, target, raw_key, kept_pins)
    print(f"  {kw}: {stats['rows']} rows -> {stats['tagged']} matched, {stats['created']} new leads, "
          f"skipped {stats['skipped']}")
    return stats


def expire_flag(conn, column, raw_key, kept_pins, min_ratio=0.5):
    kept = sorted({p for p in kept_pins if p})
    with conn.cursor() as cur:
        cur.execute(f"select count(*) from leads where {column} is true")
        flagged = cur.fetchone()[0]
        if not kept or (flagged >= 20 and len(kept) < flagged * min_ratio):
            print(f"  WARNING: {column} expiry skipped ({len(kept)} in file vs {flagged} flagged)")
            return 0
        cur.execute(
            f"update leads set {column} = null, updated_at = now() where {column} is true "
            f"and raw ? %s and coalesce(raw->%s->>'pin', '') <> all(%s::text[])", (raw_key, raw_key, kept))
        n = cur.rowcount
    print(f"  cleared {column} on {n} lead(s) no longer in the newest export")
    return n


def expire_stale_list(conn, kind, target, kw):
    raw_key = f"propstream_{target}" if kind == "tag" else f"propstream_{kw}"
    with conn.cursor() as cur:
        if kind == "tag":
            n = remove_tag(conn, target, "raw ? %(rk)s", {"rk": raw_key})
        else:
            cur.execute(f"update leads set {target} = null where {target} is true and raw ? %s", (raw_key,))
            n = cur.rowcount
    print(f"  WARNING: newest {kw} export is over {STALE_DAYS} days old -- removed it from {n} lead(s). "
          f"Re-export it from PropStream.")


def run(conn, export_dir=EXPORT_DIR, today=None):
    today = today or date.today()
    ensure_schema(conn)
    ensure_columns(conn)
    files = newest_files(export_dir)
    if not files:
        print(f"  no PropStream exports found in {export_dir}")
    for kw, kind, target in LISTS:
        path = files.get(kw)
        if not path:
            continue
        fdate = file_date(path)
        if fdate and (today - fdate).days > STALE_DAYS:
            expire_stale_list(conn, kind, target, kw)
            continue
        import_list(conn, kw, kind, target, path, today)
        conn.commit()
    rescore_all(conn)
    conn.commit()


def main():
    db_url = os.environ.get("DATABASE_URL")
    if not db_url:
        print("DATABASE_URL is not set", file=sys.stderr)
        sys.exit(1)
    conn = psycopg2.connect(db_url)
    print(f"[{datetime.now(timezone.utc).isoformat()}] Importing PropStream exports from {EXPORT_DIR}")
    run(conn)
    conn.close()
    print("Done.")


if __name__ == "__main__":
    main()
